"""Cell-level data-quality checks: non_numeric_cell, blank_cell, negative_line_item.

A text cell or an internal blank used to surface only as a reconciliation failure on the
total, which says *that* something is wrong but not *where*. These checks name the cell.
They are warnings only: every computed value must be identical with them on or off, and
no metric is blocked by them.

The workbooks are built here, in tmp_path, with plain values (no formulas), so each test
controls exactly one planted problem.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from finance_engine.config import ClientConfig, ConfigError
from finance_engine.metrics import build_metrics
from finance_engine.report import build_report
from finance_engine.sources.excel import ExcelSource
from finance_engine.validation import WARNING, validate
from tests.helpers import write_small_config

CELL_CODES = {"non_numeric_cell", "blank_cell", "negative_line_item"}

MONTHS = ["Jan 2025", "Feb 2025", "Mar 2025", "Apr 2025"]       # columns C..F
LINES = {                                                        # row: (label, values)
    7: ("Widgets", [100.0, 110.0, 120.0, 130.0]),
    8: ("Services", [50.0, 50.0, 50.0, 50.0]),
    12: ("Materials", [30.0, 30.0, 30.0, 30.0]),
    18: ("Rent", [10.0, 10.0, 10.0, 10.0]),
    20: ("Salary A", [20.0, 20.0, 20.0, 20.0]),
    21: ("Salary B", [5.0, 5.0, 5.0, 5.0]),
    23: ("Marketing spend", [40.0, 40.0, 40.0, 40.0]),
}
ON = {"data_quality": {"cell_checks": True}}


def build(path, edits: dict[tuple[int, int], object] | None = None):
    """The tiny P&L, with `edits` {(row, month_index): raw} applied to line-item cells.

    Totals are written from the unedited figures, as if the line had been changed after
    the totals were last calculated. An edit of `None` leaves the cell genuinely empty.
    """
    edits = edits or {}
    v = {row: values for row, (_, values) in LINES.items()}
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    ws["A1"] = "Income Statement"
    for i, label in enumerate(MONTHS):
        ws.cell(4, 3 + i, label)

    def put(row, label, values, label_col=2):
        ws.cell(row, label_col, label)
        for i, value in enumerate(values):
            raw = edits.get((row, i), value)
            if raw is not None:
                ws.cell(row, 3 + i, raw)

    rev = [a + b for a, b in zip(v[7], v[8])]
    cogs = v[12]
    sal = [a + b for a, b in zip(v[20], v[21])]
    opex = [r + s + m for r, s, m in zip(v[18], sal, v[23])]
    gp = [r - c for r, c in zip(rev, cogs)]
    oi = [g - o for g, o in zip(gp, opex)]

    ws["A6"] = "Income"
    for row in (7, 8):
        put(row, *LINES[row])
    put(9, "Total Income", rev, label_col=1)
    ws["A11"] = "Cost of Goods Sold"
    put(12, *LINES[12])
    put(13, "Total Cost of Goods Sold", cogs, label_col=1)
    put(15, "Gross Profit", gp)
    ws["A17"] = "Operating Expenses"
    put(18, *LINES[18])
    ws.cell(19, 2, "Salaries")
    for row in (20, 21):
        put(row, *LINES[row])
    put(22, "Total Salaries", sal)
    put(23, *LINES[23])
    put(24, "Total Operating Expenses", opex, label_col=1)
    put(26, "Operating Income", oi)
    put(28, "Net Income", oi)
    wb.save(path)
    return path


def load(tmp_path, edits=None, overrides=ON):
    cfg = {"sheets": {"income_statement": {"last_data_column": "F"}}}
    if overrides:
        cfg.update(overrides)
    config = ClientConfig.load_path(write_small_config(tmp_path, cfg))
    workbook = build(tmp_path / "tiny.xlsx", edits)
    return ExcelSource(workbook).load(config), config


def issue(report, code):
    found = [i for i in report.issues if i.code == code]
    assert len(found) == 1, f"expected one {code}, got {[i.code for i in report.issues]}"
    return found[0]


def cell_codes(report):
    return {i.code for i in report.issues} & CELL_CODES


class TestNonNumericCell:
    def test_names_the_cell_and_says_missing_not_zero(self, tmp_path):
        dataset, config = load(tmp_path, {(23, 2): "n/a"})          # Marketing, Mar
        found = issue(validate(dataset, config, "2025-03"), "non_numeric_cell")
        assert found.severity == WARNING
        assert found.cells == [{"sheet": "P&L", "cell": "E23", "label": "Marketing spend",
                                "period": "2025-03", "raw_value": "n/a"}]
        assert "P&L!E23" in found.message
        assert "treated as missing, not zero" in found.message
        assert found.to_dict()["cells"] == found.cells

    def test_affects_the_totals_for_that_period(self, tmp_path):
        dataset, config = load(tmp_path, {(23, 2): "n/a"})
        same_month = issue(validate(dataset, config, "2025-03"), "non_numeric_cell")
        for name in ("opex.total", "opex.line.marketing_spend", "opex.category.marketing",
                     "opex.total.mom_abs", "operating_income", "opex.total.ttm"):
            assert name in same_month.affects
        assert "revenue.total" not in same_month.affects
        assert "cogs.total" not in same_month.affects

        # Reporting April, the March cell feeds only April's comparisons and the TTM.
        next_month = issue(validate(dataset, config, "2025-04"), "non_numeric_cell")
        assert "opex.total" not in next_month.affects
        assert "opex.total.prior" in next_month.affects
        assert "opex.total.mom_pct" in next_month.affects

    def test_parseable_text_is_a_number_not_an_issue(self, tmp_path):
        dataset, config = load(tmp_path, {(23, 2): "40", (18, 2): "(10)", (8, 2): "$50"})
        assert "non_numeric_cell" not in cell_codes(validate(dataset, config, "2025-03"))

    def test_a_typed_dash_is_flagged(self, tmp_path):
        """The reader has no accounting-zero rule: a typed "-" is read as missing, so it
        is flagged rather than silently treated as zero. (A numeric 0 formatted to
        *display* as "-" is stored as 0 and is not text.)"""
        dataset, config = load(tmp_path, {(23, 2): "-"})
        found = issue(validate(dataset, config, "2025-03"), "non_numeric_cell")
        assert found.cells[0]["raw_value"] == "-"
        assert dataset.income_statement.by_label("Marketing spend").value("2025-03") is None


class TestBlankCell:
    def test_internal_blank(self, tmp_path):
        dataset, config = load(tmp_path, {(18, 1): None})            # Rent, Feb
        found = issue(validate(dataset, config, "2025-02"), "blank_cell")
        assert found.severity == WARNING
        assert found.cells == [{"sheet": "P&L", "cell": "D18", "label": "Rent",
                                "period": "2025-02", "raw_value": None}]
        assert "P&L!D18" in found.message
        assert "opex.total" in found.affects and "opex.line.rent" in found.affects

    def test_blank_at_either_edge_is_not_internal(self, tmp_path):
        dataset, config = load(tmp_path, {(18, 0): None, (21, 3): None})
        assert "blank_cell" not in cell_codes(validate(dataset, config, "2025-04"))

    def test_text_is_not_also_reported_as_blank(self, tmp_path):
        dataset, config = load(tmp_path, {(18, 1): "TBC"})
        assert cell_codes(validate(dataset, config, "2025-02")) == {"non_numeric_cell"}


class TestNegativeLineItem:
    def test_negative_revenue_cell(self, tmp_path):
        dataset, config = load(tmp_path, {(8, 2): -15.0})            # Services, Mar
        found = issue(validate(dataset, config, "2025-03"), "negative_line_item")
        assert found.severity == WARNING
        assert found.cells == [{"sheet": "P&L", "cell": "E8", "label": "Services",
                                "period": "2025-03", "raw_value": -15.0}]
        assert "legitimate credit" in found.message
        assert "needs confirmation" in found.message
        assert "revenue.total" in found.affects and "gross_profit" in found.affects
        assert "opex.total" not in found.affects

    def test_expected_negative_lines_are_skipped(self, tmp_path):
        overrides = dict(ON, expected_negative_lines=["Services"])
        dataset, config = load(tmp_path, {(8, 2): -15.0}, overrides)
        assert "negative_line_item" not in cell_codes(validate(dataset, config, "2025-03"))


class TestCleanWorkbook:
    def test_raises_none_of_the_three(self, tmp_path):
        dataset, config = load(tmp_path)
        for period in ("2025-01", "2025-02", "2025-03", "2025-04"):
            assert cell_codes(validate(dataset, config, period)) == set()


class TestWarningsOnly:
    EDITS = {(23, 2): "n/a", (18, 1): None, (8, 2): -15.0}

    def test_values_identical_and_nothing_blocked(self, tmp_path):
        on_dir, off_dir = tmp_path / "on", tmp_path / "off"
        on_dir.mkdir()
        off_dir.mkdir()
        ds_on, cfg_on = load(on_dir, self.EDITS)
        ds_off, cfg_off = load(off_dir, self.EDITS, {"data_quality": {"cell_checks": False}})
        doc_on = build_metrics(ds_on, cfg_on, "2025-03")
        doc_off = build_metrics(ds_off, cfg_off, "2025-03")
        assert {i["code"] for i in doc_on["data_quality_issues"]} >= CELL_CODES
        assert not {i["code"] for i in doc_off["data_quality_issues"]} & CELL_CODES

        on = {m["name"]: m for m in doc_on["metrics"]}
        off = {m["name"]: m for m in doc_off["metrics"]}
        assert on.keys() == off.keys()
        for name in on:
            assert on[name]["value"] == off[name]["value"], name
            # A warning travels with the metric ("ok" may become "warning"), but these
            # checks never block one.
            assert ((on[name]["validation_status"] == "blocked")
                    == (off[name]["validation_status"] == "blocked")), name

    def test_report_lists_the_new_warnings(self, tmp_path):
        dataset, config = load(tmp_path, self.EDITS)
        text = build_report(build_metrics(dataset, config, "2025-03"))
        data_quality = text.split("## Data quality", 1)[1]
        for code in CELL_CODES:
            assert f"`{code}`" in data_quality
        assert "P&L!E23" in data_quality


class TestSwitch:
    def test_default(self, tmp_path):
        """Text and internal blanks are on by default; negatives are opt-in."""
        dataset, config = load(tmp_path, TestWarningsOnly.EDITS, overrides=None)
        assert config.cell_checks == {"non_numeric_cell": True, "blank_cell": True,
                                      "negative_line_item": False}
        assert cell_codes(validate(dataset, config, "2025-03")) == {"non_numeric_cell",
                                                                    "blank_cell"}

    def test_one_switch_for_all(self, tmp_path):
        dataset, config = load(tmp_path, TestWarningsOnly.EDITS,
                               {"data_quality": {"cell_checks": False}})
        assert cell_codes(validate(dataset, config, "2025-03")) == set()

    def test_a_mapping_keeps_the_other_defaults(self, tmp_path):
        overrides = {"data_quality": {"cell_checks": {"negative_line_item": True}}}
        dataset, config = load(tmp_path, TestWarningsOnly.EDITS, overrides)
        assert cell_codes(validate(dataset, config, "2025-03")) == CELL_CODES

    def test_each_check_switches_separately(self, tmp_path):
        overrides = {"data_quality": {"cell_checks": {"non_numeric_cell": True,
                                                      "blank_cell": False,
                                                      "negative_line_item": False}}}
        dataset, config = load(tmp_path, TestWarningsOnly.EDITS, overrides)
        assert cell_codes(validate(dataset, config, "2025-03")) == {"non_numeric_cell"}

    def test_unknown_check_name_is_a_config_error(self, tmp_path):
        _, config = load(tmp_path, overrides={"data_quality": {"cell_checks": {"typo": 1}}})
        with pytest.raises(ConfigError):
            config.cell_checks
