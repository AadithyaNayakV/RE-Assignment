"""`init-client` (Phase 5B, task 6): draft a config by inspecting a workbook.

The generated config is judged by the engine itself: for test_clean it must give exactly
the metric values the hand-written config gives, and for Companies A and B the same. A
layout it cannot handle is reported as a draft with its problems, never forced; a workbook
that is not a monthly income statement is refused with nothing written.
"""

from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from finance_engine.cli import EXIT_DATA, EXIT_OK, EXIT_USAGE, main
from finance_engine.config import ClientConfig, Settings
from finance_engine.metrics import build_metrics
from finance_engine.sources.excel import ExcelSource

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import make_test_workbooks  # noqa: E402  (openpyxl only)


def _sha(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _values(config: ClientConfig, workbook: Path, period: str) -> dict[str, float | None]:
    doc = build_metrics(ExcelSource(workbook).load(config), config, period)
    return {m["name"]: m["value"] for m in doc["metrics"]}


def _init(root: Path, workbook: Path, client_id: str, *extra) -> int:
    return main(["init-client", "--file", str(workbook), "--client-id", client_id,
                 "--name", f"Test {client_id}", "--root", str(root), *extra])


def _assert_same_values(generated: dict, hand: dict) -> None:
    assert set(generated) == set(hand), sorted(set(generated) ^ set(hand))
    for name, value in hand.items():
        if value is None:
            assert generated[name] is None, name
        else:
            assert generated[name] == pytest.approx(value, abs=1e-9), name


@pytest.fixture
def clean_book(tmp_path):
    """test_clean's workbook, built by the same script that made the real one."""
    path = tmp_path / "in" / "test_clean.xlsx"
    path.parent.mkdir()
    make_test_workbooks.build(str(path), False, "Test Clean Co (SYNTHETIC)")
    return path


class TestMatchesTheHandWrittenConfigs:
    @pytest.mark.parametrize("period", ["2026-01", "2025-12", "2025-06"])
    def test_test_clean(self, tmp_path, clean_book, capsys, period):
        root = tmp_path / "root"
        assert _init(root, clean_book, "tc_draft") == EXIT_OK
        out = capsys.readouterr().out
        assert "Validation passed" in out and "Next: python -m finance_engine report" in out
        generated = ClientConfig.load("tc_draft", Settings.from_env(root))
        hand = ClientConfig.load("test_clean", Settings.from_env(ROOT))
        copy = root / "data" / "clients" / "tc_draft" / "source" / "test_clean.xlsx"
        _assert_same_values(_values(generated, copy, period),
                            _values(hand, clean_book, period))

    @pytest.mark.parametrize("client_id", ["company_a_normal_growth",
                                           "company_b_revenue_decline"])
    def test_companies_a_and_b(self, tmp_path, capsys, client_id):
        book = ROOT / "data" / "synthetic" / f"{client_id}.xlsx"
        before = _sha(book)
        assert _init(tmp_path, book, "drafted") == EXIT_OK
        assert _sha(book) == before                       # the original is never modified
        generated = ClientConfig.load("drafted", Settings.from_env(tmp_path))
        hand = ClientConfig.load(client_id, Settings.from_env(ROOT))
        assert generated.raw["income_statement"]["sections"][2]["subtotal_groups"] == \
            hand.raw["income_statement"]["sections"][2]["subtotal_groups"]
        copy = tmp_path / "data" / "clients" / "drafted" / "source" / book.name
        _assert_same_values(_values(generated, copy, "2026-06"),
                            _values(hand, book, "2026-06"))


class TestTheYaml:
    def test_comment_block_and_guesses(self, tmp_path, clean_book):
        _init(tmp_path, clean_book, "tc_draft")
        text = (tmp_path / "config" / "clients" / "tc_draft.yaml").read_text("utf-8")
        head = text.split("\n\n")[0]
        assert head.startswith("# Client config for tc_draft: drafted by init-client")
        for detected in ("income-statement sheet: Income Statement [high]",
                         "header row: 5 [high]", "date order: newest_first [high]",
                         "label columns: ['A', 'B'] [high]",
                         "revenue section: 'Income' (row 6) .. 'Total Income' (row 9) [high]",
                         "operating-expense categories: 4 proposed for 4 line(s) [low]",
                         "currency: USD ($) [low]"):
            assert detected in head, detected
        assert "every one a GUESS" in head and "# Needs your decision:" in head
        assert "Currency: no currency symbol" in head
        raw = yaml.safe_load(text)
        assert raw["fallback_opex_category"] == "Uncategorised (needs review)"
        assert raw["source"] == {"type": "excel", "filename": "test_clean.xlsx"}

    def test_json(self, tmp_path, clean_book, capsys):
        assert _init(tmp_path, clean_book, "tc_draft", "--json") == EXIT_OK
        payload = json.loads(capsys.readouterr().out)
        assert payload["ok"] is True and payload["status"] == "config"
        assert payload["config_path"] == "config/clients/tc_draft.yaml"
        assert payload["workbook_copy"] == "data/clients/tc_draft/source/test_clean.xlsx"
        assert payload["periods"][-1] == "2026-01"
        assert payload["next_command"] == ("python -m finance_engine report --client "
                                           "tc_draft --period 2026-01")
        assert {d["what"] for d in payload["detected"]} >= {"header row", "date order"}


class TestLayoutsItCannotHandle:
    def _statement(self, path: Path, *, total_column=False, cogs=True, orphan_total=False):
        wb = Workbook()
        ws = wb.active
        ws.title = "P&L"
        heads = ["Jan 2025", "Feb 2025", "Mar 2025"]
        if total_column:
            heads.insert(2, "Total")                  # a non-month column between months
        for j, h in enumerate(heads):
            ws.cell(1, 2 + j, h)
        rows = [("Income", None), ("Sales", 100), ("Total Income", 100)]
        if cogs:
            rows += [("Cost of Goods Sold", None), ("Hosting", 10),
                     ("Total Cost of Goods Sold", 10)]
        rows += [("Operating Expenses", None), ("Rent", 20)]
        if orphan_total:
            rows += [("Total Facilities", 20)]
        rows += [("Total Operating Expenses", 20 + (20 if orphan_total else 0))]
        for i, (label, value) in enumerate(rows, start=2):
            ws.cell(i, 1, label)
            if value is not None:
                for j in range(len(heads)):
                    ws.cell(i, 2 + j, value)
        wb.save(path)
        return path

    @pytest.mark.parametrize("kind, message", [
        ({"total_column": True}, "non-month cells between the month columns"),
        ({"cogs": False}, "No cost of goods sold total found"),
        ({"orphan_total": True}, "no 'facilities' header row sits above it"),
    ])
    def test_reported_as_a_draft(self, tmp_path, capsys, kind, message):
        book = self._statement(tmp_path / "odd.xlsx", **kind)
        before = _sha(book)
        code = _init(tmp_path / "root", book, "odd_co")
        out = capsys.readouterr().out
        assert code == EXIT_DATA
        assert "Validation FAILED" in out and message in out
        clients = tmp_path / "root" / "config" / "clients"
        assert (clients / "odd_co.draft.yaml").exists()
        assert not (clients / "odd_co.yaml").exists()
        assert message in (clients / "odd_co.draft.yaml").read_text("utf-8")
        assert _sha(book) == before

    def test_a_failed_reconciliation_makes_a_draft(self, tmp_path, capsys):
        book = self._statement(tmp_path / "off.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(book)
        wb.active["B4"] = 999                          # Total Income no longer adds up
        wb.save(book)
        assert _init(tmp_path / "root", book, "off_co") == EXIT_DATA
        assert "Reconciliation failed: revenue_items_vs_total 2025-01" in \
            capsys.readouterr().out


class TestRefusals:
    def test_a_non_financial_workbook(self, tmp_path, capsys):
        wb = Workbook()
        wb.active["A1"] = "Team lunch rota"
        wb.active["A2"] = "Monday"
        wb.active["B2"] = "Priya"
        book = tmp_path / "rota.xlsx"
        wb.save(book)
        assert _init(tmp_path / "root", book, "rota") == EXIT_USAGE
        err = capsys.readouterr().err
        assert "does not look like a monthly income statement" in err
        assert not (tmp_path / "root" / "config").exists()
        assert not (tmp_path / "root" / "data").exists()

    @pytest.mark.parametrize("name, body, message", [
        ("broken.xlsx", b"this is not a zip file", "Cannot read broken.xlsx"),
        ("table.csv", b"a,b\n1,2\n", "not an .xlsx workbook"),
    ])
    def test_an_unreadable_file(self, tmp_path, capsys, name, body, message):
        path = tmp_path / name
        path.write_bytes(body)
        assert _init(tmp_path / "root", path, "bad_co", "--json") == EXIT_USAGE
        captured = capsys.readouterr()
        assert message in captured.err
        assert json.loads(captured.out)["ok"] is False

    def test_a_missing_file(self, tmp_path, capsys):
        assert _init(tmp_path, tmp_path / "nope.xlsx", "x_co") == EXIT_USAGE
        assert "not found" in capsys.readouterr().err

    def test_a_bad_client_id(self, tmp_path, clean_book, capsys):
        assert _init(tmp_path, clean_book, "Bad Id!") == EXIT_USAGE
        assert "--client-id" in capsys.readouterr().err

    def test_never_overwrites_without_force(self, tmp_path, clean_book, capsys):
        clients = tmp_path / "config" / "clients"
        clients.mkdir(parents=True)
        (clients / "tc_draft.yaml").write_text("# mine\n", "utf-8")
        assert _init(tmp_path, clean_book, "tc_draft") == EXIT_USAGE
        assert "--force" in capsys.readouterr().err
        assert (clients / "tc_draft.yaml").read_text("utf-8") == "# mine\n"
        assert not (tmp_path / "data").exists()           # nothing copied either
        assert _init(tmp_path, clean_book, "tc_draft", "--force") == EXIT_OK
        assert (clients / "tc_draft.yaml").read_text("utf-8").startswith(
            "# Client config for tc_draft")

    def test_a_different_workbook_of_the_same_name_is_not_replaced(self, tmp_path,
                                                                   clean_book, capsys):
        dest = tmp_path / "data" / "clients" / "tc_draft" / "source"
        dest.mkdir(parents=True)
        (dest / "test_clean.xlsx").write_bytes(b"someone else's file")
        assert _init(tmp_path, clean_book, "tc_draft") == EXIT_USAGE
        assert "never overwritten" in capsys.readouterr().err
        assert (dest / "test_clean.xlsx").read_bytes() == b"someone else's file"


# ---- a workbook with a balance sheet tab ----------------------------------------------

MONTHS = 12
#: The block drafted by hand for both_tab_1 (an income statement tab and a balance sheet
#: tab, header row 5, labels A to C, months from D, Jan to Dec 2025). init-client must
#: draft the same from a workbook with that layout.
PART1_SHEET = {"name": "Balance Sheet", "header_row": 5, "label_columns": ["A", "B", "C"],
               "first_data_column": "D", "last_data_column": "O",
               "date_order": "oldest_first", "date_formats": ["%b %Y"]}
PART1_BLOCK = {
    "reported_totals": {"assets": "Total Assets",
                        "liabilities_and_equity": "Total Liabilities and Equity",
                        "liabilities": "Total Liabilities", "equity": "Total Equity"},
    "trend_lines": [{"key": "cash", "label": "Cash"},
                    {"key": "accounts_receivable", "label": "Accounts Receivable"},
                    {"key": "deferred_revenue", "label": "Deferred Revenue"}],
}


def _tab(ws, title: str, rows: list) -> None:
    """Title rows, month headers in row 5 from column D, then (label column, label, values)."""
    import datetime as dt
    ws["A1"] = "Two Tab Co (SYNTHETIC, built by the test)"
    ws["A2"] = f"{title}, monthly, Jan to Dec 2025"
    ws["A5"] = "Line"
    for m in range(MONTHS):
        ws.cell(5, 4 + m, dt.datetime(2025, m + 1, 1))
    for i, (col, label, values) in enumerate(rows, start=6):
        ws.cell(i, col, label)
        for m, v in enumerate(values or []):
            ws.cell(i, 4 + m, v)


def _two_tab_book(path: Path, *, balance: bool = True, unbalanced: str | None = None,
                  bs_rows=None) -> Path:
    """The both_tab_1 layout, with whole numbers so every total is exact.

    `bs_rows`, if given, is called with the balance-sheet series by name and returns the
    rows of the Balance Sheet tab instead of the full set.
    """
    ms = range(MONTHS)
    sales, service = [1000 + 10 * m for m in ms], [500 + 5 * m for m in ms]
    revenue = [a + b for a, b in zip(sales, service)]
    materials, labour = [200 + m for m in ms], [100] * MONTHS
    cogs = [a + b for a, b in zip(materials, labour)]
    gp = [r - c for r, c in zip(revenue, cogs)]
    salaries, rent, travel = [300] * MONTHS, [100] * MONTHS, [20 + m for m in ms]
    opex = [a + b + c for a, b, c in zip(salaries, rent, travel)]
    oi = [g - o for g, o in zip(gp, opex)]
    other = [7] * MONTHS
    wb = Workbook()
    wb.active.title = "Income Statement"
    _tab(wb.active, "Income Statement", [
        (1, "Revenue", None), (2, "Product sales", sales), (2, "Service income", service),
        (1, "Total Revenue", revenue), (1, "Cost of Goods Sold", None),
        (2, "Materials", materials), (2, "Direct labour", labour),
        (1, "Total Cost of Goods Sold", cogs), (1, "Gross Profit", gp),
        (1, "Operating Expenses", None), (2, "Salaries and wages", salaries),
        (2, "Rent", rent), (2, "Travel", travel), (1, "Total Operating Expenses", opex),
        (1, "Operating Income", oi), (2, "Other expenses", other),
        (1, "Net Income", [o - x for o, x in zip(oi, other)])])
    if balance:
        cash, ar, inv = [1000 + 50 * m for m in ms], [400 + m for m in ms], [300] * MONTHS
        current = [a + b + c for a, b, c in zip(cash, ar, inv)]
        fixed = [2000 - 10 * m for m in ms]
        assets = [a + b for a, b in zip(current, fixed)]
        ap, deferred, loan = [250] * MONTHS, [80 + m for m in ms], [900 - 10 * m for m in ms]
        liabilities = [a + b + c for a, b, c in zip(ap, deferred, loan)]
        capital = [1000] * MONTHS
        retained = [a - li - c for a, li, c in zip(assets, liabilities, capital)]
        equity = [c + r for c, r in zip(capital, retained)]
        le = [li + e for li, e in zip(liabilities, equity)]
        if unbalanced:                                  # one month's assets off by 100
            assets[int(unbalanced[5:]) - 1] += 100
        series = {"cash": cash, "assets": assets, "liabilities": liabilities,
                  "equity": equity, "le": le}
        rows = bs_rows(series) if bs_rows else [
            (1, "Assets", None), (2, "Cash", cash), (2, "Accounts Receivable", ar),
            (2, "Inventory", inv), (1, "Total Current Assets", current),
            (2, "Fixed Assets (net)", fixed), (1, "Total Assets", assets),
            (1, "Liabilities", None), (2, "Accounts Payable", ap),
            (2, "Deferred Revenue", deferred), (2, "Bank Loan", loan),
            (1, "Total Liabilities", liabilities), (1, "Equity", None),
            (2, "Share Capital", capital), (2, "Retained Earnings", retained),
            (1, "Total Equity", equity), (1, "Total Liabilities and Equity", le)]
        _tab(wb.create_sheet("Balance Sheet"), "Balance Sheet", rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _draft_cmd(root: Path, workbook: Path, client_id: str, *extra) -> int:
    return main(["init-client", "--draft", "--file", str(workbook), "--client-id", client_id,
                 "--name", f"Test {client_id}", "--root", str(root), *extra])


def _head(text: str) -> str:
    return text.split("\n\n")[0]


class TestBalanceSheet:
    def test_drafts_the_part1_block(self, tmp_path, capsys):
        book = _two_tab_book(tmp_path / "in" / "two_tab.xlsx")
        before = _sha(book)
        assert _draft_cmd(tmp_path / "root", book, "two_tab") == EXIT_OK
        out = capsys.readouterr().out
        assert "Checks: 72/72 passed; status VALID." in out
        assert "Balance sheet: tab 'Balance Sheet' of the same workbook" in out
        draft = tmp_path / "root" / "data" / "clients" / "two_tab" / "draft"
        text = (draft / "config.yaml").read_text("utf-8")
        raw = yaml.safe_load(text)
        assert raw["sheets"]["balance_sheet"] == PART1_SHEET
        assert raw["balance_sheet"] == PART1_BLOCK
        assert list(raw).index("balance_sheet") == list(raw).index("income_statement") + 1
        assert {"key": "balance_sheet_balances",
                "description": "Total Assets equals Total Liabilities and Equity"} in \
            raw["reconciliation"]["checks"]
        head = _head(text)
        assert "does not configure" not in text and "added by hand" not in text
        assert ("# Balance sheet: read from the 'Balance Sheet' tab of this workbook; the "
                "engine's balance check (Total Assets equals Total Liabilities and Equity) "
                "passed in all 12 month(s).") in head
        assert "balance-sheet label columns: ['A', 'B', 'C'] [high]" in head
        assert "'Other expenses' (row 21)" in head          # still listed for a decision
        s = json.loads((draft / "draft_summary.json").read_text("utf-8"))
        assert s["status"] == "valid" and "no_balance_sheet" not in s["warnings"]
        assert s["balance_sheet"]["sheet"] == "Balance Sheet"
        assert s["balance_sheet"]["balance_check"] == {"months_checked": 12,
                                                       "months_failed": []}
        assert _sha(book) == before

    def test_the_engine_reads_the_drafted_balance_sheet(self, tmp_path, capsys):
        book = _two_tab_book(tmp_path / "in" / "two_tab.xlsx")
        assert _init(tmp_path, book, "two_tab") == EXIT_OK
        out = capsys.readouterr().out
        assert "Validation passed: analyze 2025-12 reconciled 72/72 checks." in out
        assert ("Balance sheet: tab 'Balance Sheet' of the same workbook, header row 5, "
                "labels A, B, C, months from D.") in out
        config = ClientConfig.load("two_tab", Settings.from_env(tmp_path))
        copy = tmp_path / "data" / "clients" / "two_tab" / "source" / "two_tab.xlsx"
        doc = build_metrics(ExcelSource(copy).load(config), config, "2025-12")
        values = {m["name"]: m["value"] for m in doc["metrics"]}
        assert values["balance_sheet.cash"] == 1550                  # the Dec cell
        assert values["balance_sheet.assets"] == values["balance_sheet.liabilities_and_equity"]
        assert not [i for i in doc["data_quality_issues"] if i["code"] == "no_balance_sheet"]

    def test_without_a_balance_sheet_nothing_changes(self, tmp_path, capsys):
        both = _two_tab_book(tmp_path / "in" / "both.xlsx")
        only = _two_tab_book(tmp_path / "in2" / "both.xlsx", balance=False)
        assert _init(tmp_path / "a", both, "co") == EXIT_OK
        assert _init(tmp_path / "b", only, "co") == EXIT_OK
        capsys.readouterr()
        with_bs = (tmp_path / "a" / "config" / "clients" / "co.yaml").read_text("utf-8")
        without = (tmp_path / "b" / "config" / "clients" / "co.yaml").read_text("utf-8")
        assert "balance" not in without.casefold()
        raw_with, raw_without = yaml.safe_load(with_bs), yaml.safe_load(without)
        del raw_with["sheets"]["balance_sheet"], raw_with["balance_sheet"]
        raw_with["reconciliation"]["checks"].pop()
        assert raw_with == raw_without
        assert [c["key"] for c in raw_without["reconciliation"]["checks"]] == [
            "revenue_items_vs_total", "cogs_items_vs_total", "opex_items_vs_total",
            "opex_group_subtotals", "gross_profit", "operating_income"]
        assert _draft_cmd(tmp_path / "c", only, "co", "--json") == EXIT_OK
        s = json.loads(capsys.readouterr().out)["draft"]
        assert "balance_sheet" not in s and "balance_check" not in s["validation"]
        assert "no_balance_sheet" in s["warnings"]

    def test_a_balance_sheet_that_does_not_balance_is_rejected(self, tmp_path, capsys):
        book = _two_tab_book(tmp_path / "in" / "off.xlsx", unbalanced="2025-03")
        root = tmp_path / "root"
        assert _init(root, book, "off_co") == EXIT_DATA
        out = capsys.readouterr().out
        message = ("Balance check failed: the balance sheet does not balance in 2025-03: "
                   "'Total Assets' is 3882.0 but 'Total Liabilities and Equity' is 3782.0. "
                   "Assets must equal liabilities plus equity, so the draft is not accepted.")
        assert "Validation FAILED with 1 problem(s):" in out and message in out
        clients = root / "config" / "clients"
        assert not (clients / "off_co.yaml").exists()
        assert message in (clients / "off_co.draft.yaml").read_text("utf-8")
        assert _draft_cmd(root, book, "off_co") == EXIT_DATA
        out = capsys.readouterr().out
        assert "status INVALID" in out and message in out
        s = json.loads((root / "data" / "clients" / "off_co" / "draft" /
                        "draft_summary.json").read_text("utf-8"))
        assert s["status"] == "invalid" and s["problems"] == [message]

    def test_no_total_liabilities_and_equity_is_left_for_a_person(self, tmp_path, capsys):
        book = _two_tab_book(tmp_path / "in" / "no_le.xlsx",
                             bs_rows=lambda v: [(1, "Assets", None), (2, "Cash", v["cash"]),
                                                (1, "Total Assets", v["assets"]),
                                                (1, "Total Liabilities", v["liabilities"]),
                                                (1, "Total Equity", v["equity"])])
        assert _init(tmp_path, book, "no_le") == EXIT_OK
        capsys.readouterr()
        text = (tmp_path / "config" / "clients" / "no_le.yaml").read_text("utf-8")
        raw = yaml.safe_load(text)
        assert "balance_sheet" not in raw and "balance_sheet" not in raw["sheets"]
        assert ("The workbook has a balance sheet tab ('Balance Sheet') but it was not "
                "configured: the balance check needs exactly one liabilities and equity row "
                "and none was found") in _head(text)

    def test_two_cash_rows_are_a_decision_not_a_pick(self, tmp_path, capsys):
        book = _two_tab_book(tmp_path / "in" / "cash.xlsx",
                             bs_rows=lambda v: [(2, "Cash", v["cash"]),
                                                (1, "Total Cash", v["cash"]),
                                                (1, "Total Assets", v["assets"]),
                                                (1, "Total Liabilities and Equity", v["le"])])
        assert _init(tmp_path, book, "cash_co") == EXIT_OK
        capsys.readouterr()
        text = (tmp_path / "config" / "clients" / "cash_co.yaml").read_text("utf-8")
        raw = yaml.safe_load(text)
        assert raw["balance_sheet"]["trend_lines"] == []
        assert raw["balance_sheet"]["reported_totals"] == {
            "assets": "Total Assets", "liabilities_and_equity": "Total Liabilities and Equity"}
        head = _head(text)
        assert ("Balance sheet: several rows could be the cash trend line ('Cash' row 6, "
                "'Total Cash' row 7); none was configured.") in head
        assert "Balance sheet: no total liabilities row found" in head

    def test_two_balance_tabs_are_a_decision(self, tmp_path, capsys):
        from openpyxl import load_workbook
        book = _two_tab_book(tmp_path / "in" / "two.xlsx")
        wb = load_workbook(book)
        wb.copy_worksheet(wb["Balance Sheet"]).title = "Trial Balance"
        wb.save(book)
        assert _init(tmp_path, book, "two_bs") == EXIT_OK
        capsys.readouterr()
        text = (tmp_path / "config" / "clients" / "two_bs.yaml").read_text("utf-8")
        assert "balance_sheet" not in yaml.safe_load(text)
        assert ("Several tabs could be the balance sheet ('Balance Sheet', 'Trial "
                "Balance'); none was configured.") in _head(text)
