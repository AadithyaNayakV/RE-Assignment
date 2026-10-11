"""Proves the evals catch a regression: break something on purpose, watch it fail, revert.

    python evals/regression_demo.py data     # wrong subtotal in a TEST COPY of a workbook
    python evals/regression_demo.py code     # temporary off-by-one in the engine, reverted
    python evals/regression_demo.py both

**data**: copies case 04's workbook and config into a temporary tree and changes one
cached formula result, "Total Operations Payroll" for Jan 2026, by +1,000.00. The line
items are untouched, so the engine's own sums are unchanged. Only reconciliation can
notice. Runs Level 1 against the copy (expect FAIL), then against the real tree (expect
PASS). The repository is never modified.

**code**: edits `finance_engine/model.py` so the trailing-twelve-month window ends one
month early. That is a classic off-by-one which still produces plausible numbers. Runs
Level 1 with `--compare` against the saved baseline (expect FAIL with changed values),
restores the original file byte for byte, and runs again (expect PASS). The restore
happens in a `finally` block and is verified by hash.

Output for each step is saved under evals/runs/regression_demo_<timestamp>/.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from evals.common import new_run_dir  # noqa: E402

CASE = "04_expense_spike"
CLIENT = "company_c_expense_spike"
WORKBOOK = ROOT / "data" / "synthetic" / f"{CLIENT}.xlsx"
CONFIG = ROOT / "config" / "clients" / f"{CLIENT}.yaml"
TARGET_LABEL, TARGET_MONTH, DELTA = "Total Operations Payroll", "Jan 2026", 1000.0

MODEL = ROOT / "finance_engine" / "model.py"
GOOD = "    for back in range(months - 1, -1, -1):"
BAD = "    for back in range(months, 0, -1):          # REGRESSION DEMO: off by one"


def run(cmd: list[str], log: Path) -> int:
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True,
                          encoding="utf-8", errors="replace")
    log.write_text(proc.stdout + proc.stderr, encoding="utf-8")
    print(proc.stdout.rstrip())
    return proc.returncode


def patch_cached_value(src: Path, dest: Path, sheet: str, label: str, month: str,
                       delta: float) -> tuple[str, float, float]:
    """Copy an .xlsx, changing one formula cell's cached result and nothing else."""
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    row = next(i for i, r in enumerate(rows) if label in [str(v).strip() for v in r if v])
    col = next(c for r in rows for c, v in enumerate(r) if str(v).strip() == month)
    ref = f"{openpyxl.utils.get_column_letter(col + 1)}{row + 1}"
    old = float(rows[row][col])
    sheet_xml = _sheet_part(src, sheet)

    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == sheet_xml:
                xml = data.decode("utf-8")
                pattern = re.compile(rf'(<c r="{ref}"[^>]*>(?:(?!</c>).)*?<v>)([^<]*)(</v>)', re.S)
                xml, n = pattern.subn(lambda m: f"{m.group(1)}{old + delta}{m.group(3)}", xml)
                assert n == 1, f"cell {ref} not found in {sheet_xml}"
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    return ref, old, old + delta


def _sheet_part(path: Path, sheet: str) -> str:
    with zipfile.ZipFile(path) as z:
        wbxml = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = re.search(rf'<sheet[^>]*name="{re.escape(sheet)}"[^>]*r:id="([^"]+)"', wbxml).group(1)
    target = re.search(rf'<Relationship[^>]*Id="{rid}"[^>]*Target="([^"]+)"', rels) or \
        re.search(rf'<Relationship[^>]*Target="([^"]+)"[^>]*Id="{rid}"', rels)
    return "xl/" + target.group(1).lstrip("/").removeprefix("xl/")


def demo_data(out: Path) -> bool:
    print("\n=== DATA BREAK: wrong subtotal in a test copy ===")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "config" / "clients").mkdir(parents=True)
        (root / "data" / "synthetic").mkdir(parents=True)
        shutil.copy(CONFIG, root / "config" / "clients" / CONFIG.name)
        ref, old, new = patch_cached_value(WORKBOOK, root / "data" / "synthetic" /
                                           WORKBOOK.name, "Income Statement",
                                           TARGET_LABEL, TARGET_MONTH, DELTA)
        print(f"test copy: {TARGET_LABEL} {TARGET_MONTH} cell {ref} cached value "
              f"{old:,.2f} -> {new:,.2f} (line items untouched)")
        broken = run([sys.executable, "evals/run_engine_evals.py", CASE, "--root", str(root),
                      "--label", "regression_demo"], out / "data_break.log")
    print("\n--- same case against the real, untouched workbook ---")
    clean = run([sys.executable, "evals/run_engine_evals.py", CASE], out / "data_clean.log")
    ok = broken != 0 and clean == 0
    print(f"\nDATA BREAK {'DETECTED' if ok else 'NOT DETECTED'}: broken copy exit {broken}, "
          f"real tree exit {clean}")
    return ok


def demo_code(out: Path) -> bool:
    print("\n=== CODE BREAK: TTM window off by one (temporary) ===")
    original = MODEL.read_bytes()
    digest = hashlib.sha256(original).hexdigest()
    text = original.decode("utf-8")
    assert text.count(GOOD) == 1, "model.py changed; update the demo's GOOD line"
    try:
        MODEL.write_bytes(text.replace(GOOD, BAD).encode("utf-8"))
        print(f"edited {MODEL.relative_to(ROOT)}: {GOOD.strip()}  ->  {BAD.strip()}")
        broken = run([sys.executable, "evals/run_engine_evals.py", "--compare"],
                     out / "code_break.log")
    finally:
        MODEL.write_bytes(original)
    restored = hashlib.sha256(MODEL.read_bytes()).hexdigest() == digest
    print(f"\nreverted {MODEL.relative_to(ROOT)}: byte-identical = {restored}")
    print("\n--- after revert ---")
    clean = run([sys.executable, "evals/run_engine_evals.py", "--compare"],
                out / "code_clean.log")
    ok = broken != 0 and clean == 0 and restored
    print(f"\nCODE BREAK {'DETECTED' if ok else 'NOT DETECTED'}: broken exit {broken}, "
          f"reverted exit {clean}")
    return ok


def main(argv: list[str]) -> int:
    which = argv[0] if argv else "both"
    _, out = new_run_dir("regression_demo")
    results = []
    if which in ("data", "both"):
        results.append(demo_data(out))
    if which in ("code", "both"):
        results.append(demo_code(out))
    print(f"\nlogs: {out}")
    return 0 if all(results) else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
