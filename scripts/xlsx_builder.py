"""Writes statement-shaped .xlsx files with both formulas and cached values.

openpyxl can hold a formula or a value in a cell, not both. Real exports carry both: the
formula Excel would recalculate, and the cached result it last computed. The engine reads
the cached result and uses the formula to recognise a derived row, so a test workbook
without cached values would not exercise the real path.

This builder therefore saves the workbook with formulas, then writes the cached results
into the sheet XML directly.
"""

from __future__ import annotations

import re
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


@dataclass
class Row:
    """One written row, kept so callers can build formulas by row number."""

    row: int
    label: str
    kind: str                       # "header" | "item" | "formula"
    values: list[float | None] = field(default_factory=list)


class SheetBuilder:
    """Builds one worksheet. Rows are appended in statement order."""

    def __init__(self, workbook: "SyntheticWorkbook", name: str, *,
                 header_row: int, first_data_column: str, periods: list[str],
                 date_labels: list[str], titles: list[str],
                 section_label_column: str = "A", item_label_column: str = "B",
                 account_header: str = "Account"):
        self.wb = workbook
        self.ws = workbook.book.create_sheet(name) if workbook.book.sheetnames != [
            "Sheet"] else workbook.book["Sheet"]
        self.ws.title = name
        self.name = name
        self.header_row = header_row
        self.first_col = first_data_column
        self.periods = periods
        self.n = len(periods)
        self.section_col = section_label_column
        self.item_col = item_label_column
        self._cached: dict[int, list[float]] = {}
        self.rows: dict[str, Row] = {}
        self._next = header_row + 1

        for index, line in enumerate(titles, start=1):
            self.ws.cell(index, 1, line)
        self.ws[f"{item_label_column}{header_row}"] = account_header
        for offset, label in enumerate(date_labels):
            self.ws.cell(header_row, self._col_index(offset), label)

    # ---- geometry -----------------------------------------------------

    def _col_index(self, offset: int) -> int:
        from openpyxl.utils import column_index_from_string
        return column_index_from_string(self.first_col) + offset

    def col_letter(self, offset: int) -> str:
        return get_column_letter(self._col_index(offset))

    def skip(self, count: int = 1) -> None:
        self._next += count

    # ---- writing ------------------------------------------------------

    def header(self, label: str) -> Row:
        row = self._next
        self.ws[f"{self.section_col}{row}"] = label
        self._next += 1
        record = Row(row=row, label=label, kind="header")
        self.rows[label] = record
        return record

    def group_header(self, label: str) -> Row:
        """A subtotal-group header sits in the item column, not the section column."""
        row = self._next
        self.ws[f"{self.item_col}{row}"] = label
        self._next += 1
        record = Row(row=row, label=label, kind="header")
        self.rows[label] = record
        return record

    def item(self, label: str, values: list[float | None]) -> Row:
        if len(values) != self.n:
            raise ValueError(f"{label}: expected {self.n} values, got {len(values)}")
        row = self._next
        self.ws[f"{self.item_col}{row}"] = label
        for offset, value in enumerate(values):
            if value is None:
                continue          # a genuinely blank month, written as nothing
            self.ws.cell(row, self._col_index(offset), value)
        self._next += 1
        record = Row(row=row, label=label, kind="item", values=list(values))
        self.rows[label] = record
        return record

    def formula(self, label: str, template: str, values: list[float], *,
                in_section_column: bool = True) -> Row:
        """Write a formula row plus the cached results Excel would have left.

        `template` uses `{c}` for the current column letter, e.g. "SUM({c}8:{c}13)".
        """
        if len(values) != self.n:
            raise ValueError(f"{label}: expected {self.n} values, got {len(values)}")
        row = self._next
        column = self.section_col if in_section_column else self.item_col
        self.ws[f"{column}{row}"] = label
        for offset in range(self.n):
            letter = self.col_letter(offset)
            self.ws.cell(row, self._col_index(offset),
                         "=" + template.replace("{c}", letter))
        self._cached[row] = list(values)
        self._next += 1
        record = Row(row=row, label=label, kind="formula", values=list(values))
        self.rows[label] = record
        return record

    def cached(self) -> dict[int, list[float]]:
        return self._cached


class SyntheticWorkbook:
    def __init__(self):
        self.book = Workbook()
        self.sheets: list[SheetBuilder] = []

    def sheet(self, name: str, **kwargs) -> SheetBuilder:
        builder = SheetBuilder(self, name, **kwargs)
        self.sheets.append(builder)
        return builder

    def save(self, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        # Drop the default sheet if it was never claimed.
        if "Sheet" in self.book.sheetnames and len(self.book.sheetnames) > 1:
            del self.book["Sheet"]
        self.book.save(path)
        self._inject_cached(path)
        return path

    def _inject_cached(self, path: Path) -> None:
        sheet_order = self.book.sheetnames
        payload: dict[str, tuple[dict[int, list[float]], SheetBuilder]] = {}
        for builder in self.sheets:
            index = sheet_order.index(builder.name) + 1
            payload[f"xl/worksheets/sheet{index}.xml"] = (builder.cached(), builder)
        if not any(cached for cached, _ in payload.values()):
            return

        with zipfile.ZipFile(path) as zin:
            data = {name: zin.read(name) for name in zin.namelist()}

        for sheet_path, (cached, builder) in payload.items():
            if sheet_path not in data or not cached:
                continue
            xml = data[sheet_path].decode("utf-8")
            for row, values in cached.items():
                for offset, value in enumerate(values):
                    ref = f"{builder.col_letter(offset)}{row}"
                    pattern = re.compile(
                        rf'(<c r="{ref}"[^>]*>)(<f>.*?</f>)(?:<v\s*/>|<v>.*?</v>)?(</c>)'
                    )
                    xml, count = pattern.subn(
                        lambda m: f"{m.group(1)}{m.group(2)}<v>{value!r}</v>{m.group(3)}",
                        xml,
                    )
                    if count == 0:
                        raise AssertionError(
                            f"could not inject cached value at {builder.name}!{ref}"
                        )
            data[sheet_path] = xml.encode("utf-8")

        tmp = path.with_suffix(".tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, blob in data.items():
                zout.writestr(name, blob)
        shutil.move(str(tmp), str(path))
