"""Builds the SYNTHETIC manual-test workbooks (test_clean, test_errors, test_multi).

    python scripts/make_test_workbooks.py                 # all three clients
    python scripts/make_test_workbooks.py test_multi      # only the named client(s)
    python scripts/make_test_workbooks.py --missing       # only files that do not exist

Run from the repo root: paths are relative to the working directory, so it writes
data/clients/test_clean/source/test_clean.xlsx,
data/clients/test_errors/source/test_with_errors.xlsx, and two workbooks for test_multi,
data/clients/test_multi/source/test_multi_first_half.xlsx (Jan-Jun 2025) and
test_multi_second_half.xlsx (Jun-Dec 2025). June 2025 is in both on purpose: it is the
period that must stop the engine and ask which file to read. Their configs are
config/clients/<client>.yaml. Naming a client rebuilds only that client's files.

`ensure(root)` writes only the files that are missing. The eval runners and the pytest
session call it, so a fresh clone (where data/clients/ is git-ignored and empty) gets the
workbooks golden cases 14 to 18 read. The bytes are not reproducible (openpyxl stamps the
creation time), but the cells are: those cases record a fingerprint of the cell contents,
not of the file.
"""
import calendar, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

months=[(2025,m) for m in range(1,13)]+[(2026,1)]
n=len(months)
lab=lambda y,m: dt.date(y,m,1).strftime('%b %Y')
subs=[40000+1000*i for i in range(n)]
svc=[8000]*n; svc[-1]=9000
host=[6000+100*i for i in range(n)]
supp=[3000]*n
pay=[25000]*6+[26000]*(n-6)
mkt=[5000]*n; mkt[-1]=8000
soft=[2000]*n
rent=[3000]*n

def build(path, errors, title, keep=None):
    S=dict(subs=list(subs),svc=list(svc),host=host,supp=supp,pay=pay,mkt=list(mkt),soft=soft,rent=rent)
    bad=[0]*n
    total_off=None
    if errors:
        S['svc'][11]=0
        bad[-1]=-1500
        total_off=(7,5000)
    idx=list(range(n))
    if errors: idx=[i for i in idx if i!=5]   # drop Jun 2025
    if keep is not None: idx=[i for i in idx if i in keep]   # a subset of the months
    cols=list(reversed(idx))
    wb=Workbook(); ws=wb.active; ws.title='Income Statement'
    ws['A1']=title; ws['A1'].font=Font(bold=True,size=13)
    ws['A2']='Income Statement'; ws['A3']='SYNTHETIC test data'
    for j,i in enumerate(cols):
        c=ws.cell(row=5,column=3+j,value=lab(*months[i])); c.font=Font(bold=True); c.alignment=Alignment(horizontal='right')
    r=[6]
    def sec(label):
        ws.cell(row=r[0],column=1,value=label).font=Font(bold=True); r[0]+=1
    def item(label,fn):
        ws.cell(row=r[0],column=2,value=label)
        for j,i in enumerate(cols):
            c=ws.cell(row=r[0],column=3+j,value=fn(i)); c.number_format='#,##0;-#,##0'
        r[0]+=1; return r[0]-1
    def tot(label,fn):
        ws.cell(row=r[0],column=1,value=label).font=Font(bold=True)
        for j,i in enumerate(cols):
            c=ws.cell(row=r[0],column=3+j,value=fn(i)); c.number_format='#,##0;-#,##0'; c.font=Font(bold=True)
        r[0]+=1; return r[0]-1
    rows={}
    sec('Income')
    rows['subs']=item('Subscriptions',lambda i:S['subs'][i])
    rows['svc']=item('Services',lambda i:S['svc'][i])
    inc=lambda i:S['subs'][i]+S['svc'][i]
    rows['tinc']=tot('Total Income',lambda i:inc(i)+(total_off[1] if total_off and i==total_off[0] else 0))
    r[0]+=1
    sec('Cost of Goods Sold')
    item('Hosting',lambda i:S['host'][i]); item('Support',lambda i:S['supp'][i])
    cg=lambda i:S['host'][i]+S['supp'][i]
    tot('Total Cost of Goods Sold',cg)
    r[0]+=1
    rows['gp']=tot('Gross Profit',lambda i:inc(i)-cg(i))
    r[0]+=1
    sec('Operating Expenses')
    item('Payroll',lambda i:S['pay'][i])
    rows['mkt']=item('Marketing',lambda i:S['mkt'][i])
    item('Software',lambda i:S['soft'][i])
    rows['rent']=item('Rent',lambda i:S['rent'][i])
    if errors: rows['bad']=item('Bad Debts',lambda i:bad[i])
    ox=lambda i:S['pay'][i]+S['mkt'][i]+S['soft'][i]+S['rent'][i]+(bad[i] if errors else 0)
    rows['topex']=tot('Total Operating Expenses',ox)
    r[0]+=1
    rows['oi']=tot('Operating Income',lambda i:inc(i)-cg(i)-ox(i))
    tot('Net Income',lambda i:inc(i)-cg(i)-ox(i))
    if errors:
        col=lambda i:3+cols.index(i)
        ws.cell(row=rows['mkt'],column=col(9),value='n/a')      # Oct 2025 text
        ws.cell(row=rows['rent'],column=col(10)).value=None     # Nov 2025 blank
    ws.column_dimensions['A'].width=28; ws.column_dimensions['B'].width=16
    wb.save(path); return rows,cols

import os, sys
WORKBOOKS=[
    # client, file, title, errors, months kept (indices into `months`; None = all)
    ('test_clean','test_clean.xlsx','Test Clean Co (SYNTHETIC)',False,None),
    ('test_errors','test_with_errors.xlsx','Test Errors Co (SYNTHETIC)',True,None),
    ('test_multi','test_multi_first_half.xlsx','Test Multi Co (SYNTHETIC)',False,range(0,6)),
    ('test_multi','test_multi_second_half.xlsx','Test Multi Co (SYNTHETIC)',False,range(5,12)),
]

def main(clients=None, root=None, missing_only=False):
    """Build the workbooks under `root` (default: the working directory). Returns paths."""
    written=[]
    for client, fname, title, errors, keep in WORKBOOKS:
        if clients and client not in clients:
            continue
        folder=os.path.join(root or '.','data','clients',client,'source')
        path=os.path.join(folder,fname)
        if missing_only and os.path.exists(path):
            continue
        os.makedirs(folder,exist_ok=True)
        build(path,errors,title,None if keep is None else list(keep))
        print('written:',os.path.abspath(path),os.path.getsize(path),'bytes')
        written.append(path)
    return written

def ensure(root, clients=None):
    """Write only the workbooks that are missing under `root`. Existing files are never
    touched, so a workbook someone already has keeps its bytes."""
    return main(clients, root=str(root), missing_only=True)

if __name__ == '__main__':
    args=[a for a in sys.argv[1:] if a!='--missing']
    main(args, missing_only='--missing' in sys.argv[1:])
