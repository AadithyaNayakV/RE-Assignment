"""`init-client`: draft a client config by inspecting a workbook.

    python -m finance_engine init-client --file PATH.xlsx --client-id ID --name "Company"

No AI and no guessing presented as fact. The workbook is read with openpyxl, a config is
proposed from the layout, and every detected value is listed with how confident the
detection is. Anything ambiguous goes on a "needs your decision" list instead of being
picked silently.

The draft is then validated with the real engine: the config loader, the reader (as
`list-periods` uses it), an analysis of the latest month and its reconciliation. Only if
all of that passes is `config/clients/<ID>.yaml` written. Otherwise the file is
`config/clients/<ID>.draft.yaml` and every problem is printed. An existing config is never
overwritten without `--force`, and the original workbook is never modified: it is copied
to `data/clients/<ID>/source/`.

What a config needs (see config.py): `client_id`, `company_name`, `sheets.income_statement`
with `name`, `header_row`, `label_columns`, `first_data_column` and `date_order`, and
`income_statement.sections`, each with `key`, `header_label` and `total_label`. Every
operating-expense line gets exactly one category: the first matching rule, or the fallback.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from finance_engine import naming

CLIENT_ID = re.compile(r"[a-z0-9][a-z0-9_]{0,63}")
MIN_PERIODS = 3

#: Header formats tried, in order. A datetime cell needs none.
DATE_FORMATS = ["%b %Y", "%B %Y", "%b-%y", "%b-%Y", "%b %y", "%B-%Y", "%m/%Y", "%Y-%m",
                "%d %b %Y", "%Y-%m-%d"]

#: Section header words. The total is "Total <word>"; the header is "<word>".
SECTION_WORDS = {
    "revenue": ["income", "revenue", "revenues", "sales", "turnover"],
    "cogs": ["cost of goods sold", "cost of sales", "cogs", "cost of revenue",
             "direct costs"],
    "opex": ["operating expenses", "expenses", "overheads", "operating costs", "opex"],
}
SECTION_NAMES = {"revenue": "revenue", "cogs": "cost of goods sold",
                 "opex": "operating expenses"}

REPORTED = {
    "gross_profit": ["gross profit", "gross margin"],
    "operating_income": ["operating income", "operating profit", "net operating income",
                         "operating income (loss)", "operating profit (loss)"],
    "net_income": ["net income", "net profit", "net earnings", "net income (loss)",
                   "net profit (loss)"],
}

#: Proposed operating-expense categories, by common words. Order matters: the first rule
#: that matches wins, as in config.py ("health insurance" is a people cost, so the
#: payroll rule comes before the occupancy rule's "insurance"). Every one is a guess.
CATEGORY_RULES = [
    ("Payroll & People Costs", ["payroll", "salar", "wage", "bonus", "benefits",
                                "health insurance", "workers comp", "401k", "pension",
                                "recruit", "training", "severance", "staff"]),
    ("Commissions", ["commission"]),
    ("Contractors", ["contractor", "freelance"]),
    ("Marketing & Advertising", ["marketing", "advertis", "promotion", "events",
                                 "conference", "sponsor"]),
    ("Software, Hosting & AI", ["software", "saas", "hosting", "cloud", "ai token"]),
    ("Professional & Legal Fees", ["accounting", "legal", "consult", "advisory", "audit",
                                   "professional fees"]),
    ("Occupancy & Insurance", ["rent", "lease", "utilities", "insurance", "facilities"]),
    ("Travel, Meals & Office", ["travel", "meals", "office", "postage", "printing",
                                "telephone", "general expenses"]),
    ("Finance, Tax & Non-Cash", ["interest", "depreciation", "amortisation", "amortization",
                                 "bad debt", "tax", "bank", "merchant fees", "stock-based",
                                 "charities"]),
]
FALLBACK_CATEGORY = "Uncategorised (needs review)"

#: The thresholds every hand-written config in this repo uses. Defaults, not detections.
DEFAULT_THRESHOLDS = {
    "significant_mom_pct": 25.0, "significant_mom_abs": 10000.0,
    "significant_mom_min_abs": 1000.0, "oneoff_multiple": 3.0, "oneoff_min_abs": 5000.0,
    "oneoff_trailing_months": 6, "oneoff_window_min_abs": 25000.0, "top_movements": 5,
}
RECONCILIATION_CHECKS = [
    ("revenue_items_vs_total", "Sum of revenue line items equals the workbook's revenue total"),
    ("cogs_items_vs_total", "Sum of COGS line items equals the workbook's COGS total"),
    ("opex_items_vs_total", "Direct opex plus group subtotals equal the workbook's opex total"),
    ("opex_group_subtotals", "Each subtotal group's members equal its own subtotal row"),
    ("gross_profit", "Revenue total minus COGS total equals the reported gross profit"),
    ("operating_income", "Gross profit minus opex total equals the reported operating income"),
]


class InitError(Exception):
    """The workbook cannot be read, or is not a monthly income statement. Nothing written."""


@dataclass
class Detected:
    what: str
    value: object
    confidence: str          # high | medium | low
    how: str


@dataclass
class Draft:
    config: dict
    detected: list[Detected] = field(default_factory=list)
    decisions: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    categories: dict[str, list[str]] = field(default_factory=dict)
    #: (category name, match_contains words) in rule order, for the config.
    config_rules: list[tuple[str, list[str]]] = field(default_factory=list)
    periods: list[str] = field(default_factory=list)


@dataclass
class InitResult:
    status: str                       # config | draft
    config_path: Path
    workbook_copy: Path
    draft: Draft
    validation: dict

    @property
    def ok(self) -> bool:
        return self.status == "config"


# ---- helpers ----------------------------------------------------------------------

def _norm(text) -> str:
    s = str(text or "").replace("–", "-").replace("—", "-").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip().casefold()


def _number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _period(raw) -> tuple[str, str] | None:
    """(YYYY-MM, format used) for a header cell, or None."""
    if isinstance(raw, (dt.datetime, dt.date)):
        return f"{raw.year:04d}-{raw.month:02d}", "datetime cell"
    if not isinstance(raw, str) or not raw.strip():
        return None
    text = " ".join(raw.replace("\xa0", " ").split())
    for fmt in DATE_FORMATS:
        try:
            parsed = dt.datetime.strptime(text, fmt)
        except ValueError:
            continue
        return f"{parsed.year:04d}-{parsed.month:02d}", fmt
    return None


def _sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _read(path: Path) -> dict[str, list[list]]:
    """Every sheet's cached values. Read-only; the file is never written."""
    if not path.exists():
        raise InitError(f"Workbook not found: {path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise InitError(f"{path.name} is not an .xlsx workbook; init-client reads .xlsx only.")
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise InitError(f"Cannot read {path.name} as an .xlsx workbook: "
                        f"{' '.join(str(exc).split())[:200]}") from None
    try:
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)]
                for ws in wb.worksheets}
    finally:
        wb.close()


# ---- detection ----------------------------------------------------------------------

def _header_row(rows: list[list]) -> tuple[int, list[tuple[int, str, str]]] | None:
    """The first row with at least MIN_PERIODS month labels: (row number, [(col, key, fmt)])."""
    for i, row in enumerate(rows[:60], start=1):
        found = [(c, *p) for c, v in enumerate(row, start=1) if (p := _period(v))]
        if len(found) >= MIN_PERIODS:
            return i, found
    return None


def _score_sheet(name: str, rows: list[list]) -> tuple[int, dict | None]:
    header = _header_row(rows)
    if header is None:
        return 0, None
    labels = {_norm(v) for r in rows for v in r[:4] if isinstance(v, str)}
    sections = sum(any(f"total {w}" in labels for w in words)
                   for words in SECTION_WORDS.values())
    named = any(w in _norm(name) for w in ("income", "profit", "p&l", "p & l", "pl", "loss"))
    return 10 * sections + (3 if named else 0) + 1, {"header": header}


def detect(path: Path, client_id: str, company_name: str) -> Draft:
    sheets = _read(path)
    candidates = []
    for name, rows in sheets.items():
        score, info = _score_sheet(name, rows)
        if info and "balance" not in _norm(name):
            candidates.append((score, info, name))
    candidates.sort(key=lambda c: -c[0])            # stable: workbook order breaks ties
    if not candidates:
        raise InitError(
            f"{path.name}: no sheet has a row with at least {MIN_PERIODS} month labels "
            f"(tried {', '.join(DATE_FORMATS)} and date cells). This does not look like a "
            f"monthly income statement, so no config was drafted.")
    d = Draft(config={})
    best_score, info, sheet = candidates[0]
    ties = [n for s, _, n in candidates[1:] if s == best_score]
    d.detected.append(Detected(
        "income-statement sheet", sheet, "high" if best_score >= 31 and not ties else
        ("medium" if best_score >= 11 else "low"),
        "has month headers" + (" and all three section totals" if best_score >= 31 else
                               " but not all three section totals")))
    if ties:
        d.decisions.append(f"Sheets {[sheet, *ties]} look equally like an income statement; "
                           f"{sheet!r} was used. Change sheets.income_statement.name if wrong.")
    others = [n for n in sheets if n != sheet]
    if any("balance" in _norm(n) for n in others):
        d.decisions.append(
            f"The workbook has a balance sheet ({', '.join(n for n in others if 'balance' in _norm(n))}). "
            f"init-client does not configure it; balance-sheet trends stay unavailable until "
            f"a balance_sheet block is added by hand.")
    rows = sheets[sheet]
    header_no, periods = info["header"]
    d.detected.append(Detected("header row", header_no, "high",
                               f"first row with {len(periods)} parseable month labels"))

    cols = [c for c, _, _ in periods]
    keys = [k for _, k, _ in periods]
    formats = sorted({f for _, _, f in periods if f != "datetime cell"})
    first, last = min(cols), max(cols)
    header_cells = rows[header_no - 1]
    stray = [get_column_letter(c) for c in range(first, last + 1)
             if c not in cols and c - 1 < len(header_cells)
             and header_cells[c - 1] not in (None, "")]
    trailing = [get_column_letter(c) for c in range(last + 1, len(header_cells) + 1)
                if header_cells[c - 1] not in (None, "")]
    if stray:
        d.problems.append(f"Header row {header_no} has non-month cells between the month "
                          f"columns ({', '.join(stray)}). The engine reads every column "
                          f"between the first and last month and cannot skip them.")
    if len(set(keys)) != len(keys):
        dupes = sorted({k for k in keys if keys.count(k) > 1})
        d.problems.append(f"Duplicate month columns: {', '.join(dupes)}.")
    ordinals = [int(k[:4]) * 12 + int(k[5:]) for k in keys]
    if all(a < b for a, b in zip(ordinals, ordinals[1:])):
        order = "oldest_first"
    elif all(a > b for a, b in zip(ordinals, ordinals[1:])):
        order = "newest_first"
    else:
        order = "newest_first"
        d.problems.append("The month columns are not in date order either way; the engine "
                          "needs them strictly oldest-first or newest-first.")
    d.periods = sorted(set(keys))
    d.detected.append(Detected("first data column", get_column_letter(first), "high",
                               "the first month column"))
    d.detected.append(Detected("last data column", get_column_letter(last), "high",
                               "the last month column" + (
                                   f"; columns after it ({', '.join(trailing)}) are not "
                                   f"months and are excluded" if trailing else "")))
    d.detected.append(Detected("date order", order, "high" if not d.problems else "low",
                               f"{keys[0]} in the first column, {keys[-1]} in the last"))
    d.detected.append(Detected("date format(s)", formats or ["datetime cells"],
                               "high", "every month header parsed with these"))

    # Label columns: text below the header, left of the first month column.
    label_cols = [c for c in range(1, first)
                  if any(isinstance(r[c - 1], str) and r[c - 1].strip()
                         for r in rows[header_no:] if c - 1 < len(r))]
    if not label_cols:
        raise InitError(f"{path.name}!{sheet}: no text labels to the left of the month "
                        f"columns, so line items cannot be named.")
    d.detected.append(Detected("label columns", [get_column_letter(c) for c in label_cols],
                               "high", "columns left of the months that hold text"))

    lines = []                                   # (row, label, has_numbers)
    for i, r in enumerate(rows[header_no:], start=header_no + 1):
        label = next((" ".join(str(r[c - 1]).split()) for c in label_cols
                      if c - 1 < len(r) and r[c - 1] not in (None, "")
                      and str(r[c - 1]).strip()), None)
        nums = any(_number(r[c - 1]) for c in cols if c - 1 < len(r))
        if label:
            lines.append((i, label, nums))
    by_norm: dict[str, list[int]] = {}
    for n, label, _ in lines:
        by_norm.setdefault(_norm(label), []).append(n)

    sections, bounds = _sections(d, lines, by_norm)
    reported = _reported(d, lines, by_norm)
    _categories(d, lines, bounds)
    _outside_rows(d, lines, bounds, reported)
    _negatives_and_units(d, rows, lines, bounds, cols, header_no)

    currency, symbol = _currency(path, sheet)
    if currency is None:
        currency, symbol = "USD", "$"
        d.decisions.append("Currency: no currency symbol in the cell formats. USD and '$' "
                           "were assumed; correct currency and currency_symbol if wrong.")
        d.detected.append(Detected("currency", "USD ($)", "low",
                                   "assumed: the cell formats carry no symbol"))
    else:
        d.detected.append(Detected("currency", f"{currency} ({symbol})", "medium",
                                   "from the cell number formats"))

    d.config = {
        "client_id": client_id,
        "company_name": company_name,
        "currency": currency,
        "currency_symbol": symbol,
        "units_scale": 1,
        "source": {"type": "excel", "filename": path.name},
        "sheets": {"income_statement": {
            "name": sheet, "header_row": header_no,
            "label_columns": [get_column_letter(c) for c in label_cols],
            "first_data_column": get_column_letter(first),
            "last_data_column": get_column_letter(last),
            "date_order": order,
            "date_formats": formats or ["%b %Y"],
        }},
        "income_statement": {"sections": sections, "reported_lines": reported},
        "opex_categories": [{"name": name, "match_contains": needles}
                            for name, needles in d.config_rules],
        # A clearly named fallback, so an unmatched line is visibly unreviewed.
        "fallback_opex_category": FALLBACK_CATEGORY,
        "reconciliation": {"tolerance_abs": 0.01,
                           "checks": [{"key": k, "description": v}
                                      for k, v in RECONCILIATION_CHECKS]},
        "classification_notes": [],
        "thresholds": dict(DEFAULT_THRESHOLDS),
        "ttm": {"months": 12},
    }
    return d


def _sections(d: Draft, lines, by_norm) -> tuple[list[dict], dict]:
    """Section header and total labels, and each section's row bounds."""
    sections, bounds, floor = [], {}, 0
    for key in ("revenue", "cogs", "opex"):
        totals = [(n, label, w) for w in SECTION_WORDS[key]
                  for n, label, _ in lines if _norm(label) == f"total {w}" and n > floor]
        totals.sort()
        if not totals:
            d.problems.append(
                f"No {SECTION_NAMES[key]} total found (looked for "
                f"{', '.join(repr('Total ' + w.title()) for w in SECTION_WORDS[key])}"
                + (f" below row {floor}" if floor else "") + ")."
                + (" Gross profit is revenue minus cost of goods sold, so the engine cannot "
                   "analyse a workbook without one." if key == "cogs" else ""))
            continue
        t_row, t_label, word = totals[0]
        if len(totals) > 1:
            d.decisions.append(
                f"{SECTION_NAMES[key].capitalize()}: several total rows could be the section "
                f"total ({', '.join(f'{l!r} row {n}' for n, l, _ in totals)}); the first, "
                f"{t_label!r}, was used.")
        if len(by_norm.get(_norm(t_label), [])) > 1:
            d.problems.append(f"{t_label!r} appears on rows {by_norm[_norm(t_label)]}; the "
                              f"engine refuses an ambiguous anchor.")
        above = [(n, label, nums) for n, label, nums in lines if floor < n < t_row]
        exact = [(n, label) for n, label, nums in above if _norm(label) == word and not nums]
        loose = [(n, label) for n, label, nums in above if not nums
                 and any(w in _norm(label) for w in SECTION_WORDS[key])]
        if exact:
            h_row, h_label = exact[-1]
            confidence, how = "high", f"{h_label!r} heads {t_label!r}"
        elif loose:
            h_row, h_label = loose[-1]
            confidence, how = "medium", f"nearest heading row above {t_label!r} with a {SECTION_NAMES[key]} word"
            d.decisions.append(f"{SECTION_NAMES[key].capitalize()} header {h_label!r} (row "
                               f"{h_row}) was matched loosely to {t_label!r}; confirm it.")
        else:
            d.problems.append(f"No header row found above {t_label!r} (row {t_row}).")
            continue
        groups = _groups(d, key, lines, h_row, t_row)
        sec = {"key": key, "header_label": h_label, "total_label": t_label}
        if key == "opex" or groups:
            sec["subtotal_groups"] = groups
        sections.append(sec)
        bounds[key] = (h_row, t_row)
        floor = t_row
        d.detected.append(Detected(f"{SECTION_NAMES[key]} section",
                                   f"{h_label!r} (row {h_row}) .. {t_label!r} (row {t_row})",
                                   confidence, how))
    return sections, bounds


def _groups(d: Draft, key: str, lines, top: int, bottom: int) -> list[dict]:
    """Subtotal groups inside a section: a "Total X" row with an "X" header above it."""
    groups, inside = [], [(n, label, nums) for n, label, nums in lines if top < n < bottom]
    for n, label, nums in inside:
        m = re.match(r"total\s+(.+)", _norm(label))
        if not m:
            continue
        name = m.group(1)
        header = [(hn, hl) for hn, hl, hnums in inside
                  if hn < n and _norm(hl) == name and not hnums]
        if not header:
            d.problems.append(
                f"{label!r} (row {n}) looks like a subtotal inside "
                f"{SECTION_NAMES[key]} but no {name!r} header row sits above it. Without a "
                f"subtotal group in the config its members would be counted twice.")
            continue
        hn, hl = header[-1]
        groups.append({"key": naming.slug(hl), "name": hl, "header_label": hl,
                       "total_label": label})
        d.detected.append(Detected(f"subtotal group in {SECTION_NAMES[key]}",
                                   f"{hl!r} (row {hn}) .. {label!r} (row {n})", "high",
                                   "a 'Total X' row with an 'X' heading above it"))
    return groups


def _reported(d: Draft, lines, by_norm) -> dict:
    reported = {}
    for key, labels in REPORTED.items():
        found = [(n, label) for n, label, nums in lines if nums and _norm(label) in labels]
        if not found:
            d.decisions.append(f"No {key.replace('_', ' ')} row found (looked for "
                               f"{', '.join(repr(x) for x in labels)}). The engine still "
                               f"computes it; only the reconciliation against the workbook's "
                               f"own row is skipped.")
            continue
        if len({_norm(l) for _, l in found}) > 1 or len(found) > 1:
            d.problems.append(f"Several rows could be {key.replace('_', ' ')}: "
                              f"{', '.join(f'{l!r} row {n}' for n, l in found)}.")
            continue
        n, label = found[0]
        reported[key] = label
        d.detected.append(Detected(f"reported {key.replace('_', ' ')}", f"{label!r} (row {n})",
                                   "high", "exact label match"))
    return reported


def _categories(d: Draft, lines, bounds) -> None:
    """Propose opex categories by common words. Every line lands in exactly one."""
    if "opex" not in bounds:
        return
    top, bottom = bounds["opex"]
    items = [(n, label) for n, label, nums in lines
             if top < n < bottom and nums and not _norm(label).startswith("total ")]
    used: dict[str, list[str]] = {}
    for n, label in items:
        text = _norm(label)
        hits = [(name, needle) for name, needles in CATEGORY_RULES for needle in needles
                if needle in text]
        if not hits:
            d.categories.setdefault(FALLBACK_CATEGORY, []).append(label)
            continue
        name, needle = hits[0]
        d.categories.setdefault(name, []).append(label)
        used.setdefault(name, [])
        if needle not in used[name]:
            used[name].append(needle)
        rivals = sorted({h for h, _ in hits if h != name})
        if rivals:
            d.decisions.append(f"Category of {label!r}: matched {name!r} and also "
                               f"{', '.join(map(repr, rivals))}; the first rule wins. Confirm.")
    # Keep CATEGORY_RULES order so first-match-wins gives the assignments shown above.
    d.config_rules = [(name, used[name]) for name, _ in CATEGORY_RULES if name in used]
    if d.categories.get(FALLBACK_CATEGORY):
        d.decisions.append(
            f"{len(d.categories[FALLBACK_CATEGORY])} operating-expense line(s) matched no "
            f"category word and sit in {FALLBACK_CATEGORY!r}: "
            f"{', '.join(map(repr, d.categories[FALLBACK_CATEGORY]))}. Add a rule or move them.")
    d.detected.append(Detected("operating-expense categories",
                               f"{len(d.config_rules)} proposed for {len(items)} line(s)",
                               "low", "match_contains on common words; every one a guess"))


def _outside_rows(d: Draft, lines, bounds, reported) -> None:
    inside = lambda n: any(top <= n <= bottom for top, bottom in bounds.values())
    rep = {_norm(v) for v in reported.values()}
    stray = [f"{label!r} (row {n})" for n, label, nums in lines
             if nums and not inside(n) and _norm(label) not in rep]
    if stray:
        d.decisions.append(f"Rows with numbers outside the detected sections are ignored by "
                           f"the engine: {', '.join(stray[:8])}"
                           f"{' ...' if len(stray) > 8 else ''}. Confirm none belongs in a "
                           f"section.")


def _negatives_and_units(d: Draft, rows, lines, bounds, cols, header_no) -> None:
    for key in ("revenue", "cogs", "opex"):
        if key not in bounds:
            continue
        top, bottom = bounds[key]
        for n, label, nums in lines:
            if not (top < n < bottom and nums) or _norm(label).startswith("total "):
                continue
            values = [rows[n - 1][c - 1] for c in cols if c - 1 < len(rows[n - 1])]
            negative = [v for v in values if _number(v) and v < 0]
            if negative and len(negative) * 2 >= len([v for v in values if _number(v)]):
                d.decisions.append(
                    f"{label!r} is negative in {len(negative)} month(s). If that is its "
                    f"nature (interest income, a tax benefit), list it under "
                    f"expected_negative_lines.")
    titles = " ".join(str(v) for r in rows[:header_no] for v in r if isinstance(v, str))
    if re.search(r"\b(thousand|thousands|million|millions|lakh|crore)\b", titles, re.I):
        d.decisions.append("The sheet heading mentions a unit scale; units_scale is 1 "
                           "(plain units). Set it if the figures are in thousands or millions.")


def _currency(path: Path, sheet: str) -> tuple[str | None, str | None]:
    """A currency from the cells' number formats, if they carry one."""
    try:
        wb = load_workbook(path, read_only=True)
        formats = [c.number_format for r in wb[sheet].iter_rows(max_row=60) for c in r
                   if getattr(c, "number_format", None)]
        wb.close()
    except Exception:            # a currency hint is optional; never fail on it
        return None, None
    text = " ".join(formats)
    for sym, code in (("$", "USD"), ("£", "GBP"), ("€", "EUR"), ("₹", "INR")):
        if sym in text:
            return code, sym
    return None, None


# ---- validation and writing ------------------------------------------------------------

def _validate(raw: dict, copy: Path, root: Path) -> tuple[list[str], dict]:
    """Load with the real config loader, read the periods, analyse the latest month."""
    from finance_engine.config import ClientConfig, ConfigError
    from finance_engine.metrics import build_metrics
    from finance_engine.model import LayoutError
    from finance_engine.sources.excel import ExcelSource

    problems: list[str] = []
    info: dict = {}
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / f"{raw['client_id']}.yaml"
        path.write_text(yaml.safe_dump(raw, sort_keys=False, allow_unicode=True),
                        encoding="utf-8")
        try:
            config = ClientConfig.load_path(path)
        except ConfigError as exc:
            return [f"The config loader rejected the draft: {exc}"], info
    try:
        dataset = ExcelSource(copy).load(config)
    except (LayoutError, KeyError, ValueError) as exc:
        return [f"list-periods failed: {exc}"], info
    keys = dataset.period_keys()
    info["periods"] = keys
    if not keys:
        return ["list-periods found no months."], info
    latest = keys[-1]
    try:
        doc = build_metrics(dataset, config, latest, project_root=root)
    except (LayoutError, KeyError, ValueError) as exc:
        return [f"analyze {latest} failed: {exc}"], info
    recon = doc["reconciliation"]
    info.update(latest=latest, checks_run=recon["checks_run"],
                checks_failed=recon["checks_failed"])
    for f in recon["failures"][:6]:
        problems.append(f"Reconciliation failed: {f['key']} {f['period']} (computed "
                        f"{f['computed']}, workbook {f['reported']}). The detected section "
                        f"or group boundaries are probably wrong.")
    if recon["checks_failed"] > 6:
        problems.append(f"... and {recon['checks_failed'] - 6} more reconciliation failures.")
    for issue in doc["data_quality_issues"]:
        if issue["severity"] == "error" and issue["code"] != "ttm_window_incomplete":
            problems.append(f"analyze {latest}: {issue['code']}: {issue['message']}")
    info["warnings"] = [i["code"] for i in doc["data_quality_issues"]
                        if i["severity"] == "warning"]
    return problems, info


def _comment_block(draft: Draft, source: Path, status: str, validation: dict) -> str:
    today = dt.date.today().isoformat()
    lines = [f"# Client config for {draft.config['client_id']}: drafted by init-client on "
             f"{today}",
             f"# from {source.name}. Every value below was detected from the workbook's "
             f"layout, not confirmed.",
             "#"]
    if status == "config":
        lines.append(f"# Validation: PASSED. The real config loader accepted it, list-periods "
                     f"found {len(validation.get('periods', []))} month(s), and analyze "
                     f"{validation.get('latest')} reconciled "
                     f"{validation.get('checks_run', 0) - validation.get('checks_failed', 0)}"
                     f"/{validation.get('checks_run', 0)} checks.")
    else:
        lines.append("# Validation: FAILED. This is a DRAFT and the engine will not load it "
                     "under --client.")
        lines.append("# Problems:")
        lines += [f"#   - {p}" for p in draft.problems]
    lines += ["#", "# Detected (confidence):"]
    for x in draft.detected:
        lines.append(f"#   {x.what}: {x.value} [{x.confidence}] ({x.how})")
    if draft.categories:
        lines += ["#", "# Proposed operating-expense categories (every one a GUESS):"]
        for name, labels in draft.categories.items():
            lines.append(f"#   {name}: {', '.join(labels)}")
    lines += ["#", "# Needs your decision:"]
    decisions = draft.decisions + [
        "Thresholds and the TTM length are this repo's defaults, not detected; tune them "
        "per client.",
        "Classification notes are empty; add anything a reviewer must know about how this "
        "company presents its statements."]
    lines += [f"#   - {x}" for x in decisions]
    return "\n".join(lines) + "\n\n"


def init_client(source: Path, client_id: str, company_name: str, root: Path,
                config_dir: Path, data_dir: Path, force: bool = False) -> InitResult:
    source = Path(source)
    if not CLIENT_ID.fullmatch(client_id):
        raise InitError(f"--client-id {client_id!r}: use lowercase letters, digits and "
                        f"underscores, starting with a letter or digit.")
    if not company_name.strip():
        raise InitError("--name must not be empty.")
    final, draft_path = config_dir / f"{client_id}.yaml", config_dir / f"{client_id}.draft.yaml"
    existing = [p for p in (final, draft_path) if p.exists()]
    if existing and not force:
        raise InitError(f"{', '.join(str(p) for p in existing)} already exists. Re-run with "
                        f"--force to replace it; nothing was changed.")

    draft = detect(source, client_id, company_name.strip())   # raises before any write

    dest_dir = data_dir / client_id / "source"
    dest = dest_dir / source.name
    if dest.exists() and dest.resolve() != source.resolve() and _sha(dest) != _sha(source):
        raise InitError(f"{dest} already exists and differs from {source}. Rename one; a "
                        f"client's source workbook is never overwritten.")
    dest_dir.mkdir(parents=True, exist_ok=True)
    if not dest.exists():
        shutil.copy2(source, dest)

    problems, validation = (draft.problems, {}) if draft.problems else \
        _validate(draft.config, dest, root)
    draft.problems = list(problems)
    status = "draft" if draft.problems else "config"
    path = final if status == "config" else draft_path
    config_dir.mkdir(parents=True, exist_ok=True)
    body = yaml.safe_dump(draft.config, sort_keys=False, allow_unicode=True)
    path.write_text(_comment_block(draft, source, status, validation) + body, encoding="utf-8")
    if status == "config" and draft_path.exists() and force:
        draft_path.unlink()          # a stale draft would contradict the new config
    return InitResult(status, path, dest, draft, validation)


def summary_lines(result: InitResult, root: Path) -> list[str]:
    d, cfg = result.draft, result.draft.config
    sheet = cfg["sheets"]["income_statement"]
    rel = lambda p: _rel(p, root)
    sections = cfg["income_statement"]["sections"]
    groups = sum(len(s.get("subtotal_groups") or []) for s in sections)
    lines = [
        f"Detected in {result.workbook_copy.name}: sheet {sheet['name']!r}, header row "
        f"{sheet['header_row']}, {len(d.periods)} month(s)"
        + (f" ({d.periods[0]} to {d.periods[-1]}, {sheet['date_order'].replace('_', ' ')})"
           if d.periods else "") + ".",
        f"Sections: " + ("; ".join(f"{s['header_label']!r} .. {s['total_label']!r}"
                                   for s in sections) or "none found")
        + f". Subtotal groups: {groups}. Reported lines: "
        + (", ".join(cfg['income_statement']['reported_lines'].values()) or "none") + ".",
        f"Operating-expense categories proposed: {len(cfg['opex_categories'])} (guesses)"
        + (f"; {len(d.categories.get(FALLBACK_CATEGORY, []))} line(s) uncategorised"
           if d.categories.get(FALLBACK_CATEGORY) else "") + ".",
        f"Workbook copied to {rel(result.workbook_copy)} (the original is untouched).",
    ]
    if result.ok:
        v = result.validation
        lines.append(f"Validation passed: analyze {v.get('latest')} reconciled "
                     f"{v.get('checks_run', 0) - v.get('checks_failed', 0)}/"
                     f"{v.get('checks_run', 0)} checks.")
        lines.append(f"Wrote {rel(result.config_path)}. {len(d.decisions)} item(s) need your "
                     f"decision; they are listed at the top of the file.")
        if d.periods:
            lines.append(f"Next: python -m finance_engine report --client "
                         f"{cfg['client_id']} --period {d.periods[-1]}")
    else:
        lines.append(f"Validation FAILED with {len(d.problems)} problem(s):")
        lines += [f"  - {p}" for p in d.problems]
        lines.append(f"Wrote {rel(result.config_path)} (a draft; --client will not load it).")
        lines.append(f"Next: fix the problems in the draft, rename it to "
                     f"{cfg['client_id']}.yaml, then run python -m finance_engine "
                     f"list-periods --client {cfg['client_id']}")
    return lines


def next_command(result: InitResult) -> str:
    cfg = result.draft.config
    if result.ok and result.draft.periods:
        return (f"python -m finance_engine report --client {cfg['client_id']} "
                f"--period {result.draft.periods[-1]}")
    return f"python -m finance_engine list-periods --client {cfg['client_id']}"


def _rel(path: Path, root: Path) -> str:
    try:
        return Path(path).resolve().relative_to(Path(root).resolve()).as_posix()
    except ValueError:
        return Path(path).as_posix()
