"""Excel source: reads .xlsx with openpyxl, twice.

The workbook is opened once with `data_only=False` to see formulas and once with
`data_only=True` to see the values Excel last cached. Both matter: the formulas tell us
which rows the workbook derives itself (so we can reconcile against them), and the cached
values are the only numbers available without an Excel engine.

The source file is opened read-only and never written.
"""

from __future__ import annotations

import datetime as dt
import hashlib
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from finance_engine.config import ClientConfig
from finance_engine.model import Cell, Dataset, LayoutError, LineItem, Period, Sheet, _norm
from finance_engine.sources.base import DataSource


class ExcelSource(DataSource):
    source_type = "excel"

    def __init__(self, path: Path):
        self.path = Path(path)
        if not self.path.exists():
            raise LayoutError(f"Workbook not found: {self.path}")
        self._sha256: str | None = None

    # ---- DataSource ---------------------------------------------------

    def load(self, config: ClientConfig) -> Dataset:
        wb_formula = load_workbook(self.path, data_only=False, read_only=False)
        wb_values = load_workbook(self.path, data_only=True, read_only=False)
        try:
            is_spec = config.sheet_spec("income_statement")
            income = self._read_sheet(wb_formula, wb_values, is_spec, config)
            self._apply_income_statement_structure(income, config)

            balance: Sheet | None = None
            bs_spec = config.sheet_spec("balance_sheet")
            if bs_spec and bs_spec.get("name") in wb_formula.sheetnames:
                balance = self._read_sheet(wb_formula, wb_values, bs_spec, config)

            return Dataset(
                client_id=config.client_id,
                company_name=config.company_name,
                currency=config.currency,
                source_type=self.source_type,
                source_path=str(self.path),
                source_sha256=self.sha256(),
                income_statement=income,
                balance_sheet=balance,
            )
        finally:
            wb_formula.close()
            wb_values.close()

    def describe(self) -> dict:
        return {
            "type": self.source_type,
            "path": str(self.path),
            "filename": self.path.name,
            "sha256": self.sha256(),
            "size_bytes": self.path.stat().st_size,
            "modified_utc": dt.datetime.fromtimestamp(
                self.path.stat().st_mtime, tz=dt.timezone.utc
            ).isoformat(),
        }

    def sha256(self) -> str:
        if self._sha256 is None:
            h = hashlib.sha256()
            with self.path.open("rb") as fh:
                for chunk in iter(lambda: fh.read(65536), b""):
                    h.update(chunk)
            self._sha256 = h.hexdigest()
        return self._sha256

    # ---- parsing ------------------------------------------------------

    def _read_sheet(self, wb_formula, wb_values, spec: dict, config: ClientConfig) -> Sheet:
        name = spec["name"]
        if name not in wb_formula.sheetnames:
            raise LayoutError(
                f"{self.path.name}: expected a sheet named {name!r}; "
                f"found {wb_formula.sheetnames}"
            )
        sf, sv = wb_formula[name], wb_values[name]

        header_row = int(spec["header_row"])
        first_col = column_index_from_string(str(spec["first_data_column"]))
        last_col = (
            column_index_from_string(str(spec["last_data_column"]))
            if spec.get("last_data_column")
            else sf.max_column
        )
        label_cols = [column_index_from_string(str(c)) for c in spec["label_columns"]]
        date_formats = list(spec.get("date_formats") or ["%b %Y"])

        periods, period_cols = self._read_periods(
            sv, name, header_row, first_col, last_col, date_formats, spec["date_order"]
        )

        titles: list[str] = []
        for row in range(1, header_row):
            for col in range(1, min(sf.max_column, 6) + 1):
                raw = sf.cell(row, col).value
                if isinstance(raw, str) and raw.strip():
                    titles.append(" ".join(raw.split()))

        sheet = Sheet(name=name, periods=periods, titles=titles,
                      period_columns={p.key: get_column_letter(c) for p, c in period_cols})
        for row in range(header_row + 1, sf.max_row + 1):
            label = self._row_label(sf, row, label_cols)
            cells: dict[str, Cell] = {}
            for period, col_idx in period_cols:
                cell_f = sf.cell(row, col_idx)
                cell_v = sv.cell(row, col_idx)
                raw_f, raw_v = cell_f.value, cell_v.value
                is_formula = isinstance(raw_f, str) and raw_f.startswith("=")
                value, unreadable = _coerce(raw_v)
                if raw_f is None and raw_v is None:
                    continue
                cells[period.key] = Cell(
                    sheet=name,
                    row=row,
                    column=cell_f.column_letter,
                    label=label or "",
                    period=period.key,
                    value=value,
                    is_formula=is_formula,
                    formula=raw_f if is_formula else None,
                    # A formula whose cached value is absent or non-numeric cannot be
                    # evaluated without Excel. Flagged, never guessed at.
                    unreadable_formula=is_formula and (raw_v is None or unreadable),
                    unreadable_text=(raw_v.strip() if not is_formula and unreadable
                                     and isinstance(raw_v, str) else None),
                )
            if not label and not cells:
                continue
            has_numbers = any(c.value is not None for c in cells.values())
            sheet.line_items.append(
                LineItem(
                    sheet=name,
                    row=row,
                    label=label or "",
                    role="item" if has_numbers else "header",
                    cells=cells,
                )
            )
        return sheet

    def _read_periods(self, sv, sheet_name, header_row, first_col, last_col,
                      date_formats, date_order):
        """Parse the period header into `Period`s, always returned oldest-first."""
        found: list[tuple[Period, int]] = []
        for col in range(first_col, last_col + 1):
            raw = sv.cell(header_row, col).value
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                continue
            period = _parse_period(raw, date_formats)
            if period is None:
                raise LayoutError(
                    f"{sheet_name}!{sv.cell(header_row, col).coordinate}: cannot read "
                    f"{raw!r} as a period using formats {date_formats}. "
                    f"Fix the client config rather than letting the engine guess."
                )
            found.append((period, col))

        if not found:
            raise LayoutError(
                f"{sheet_name}: no period headers found in row {header_row}, "
                f"columns {first_col}-{last_col}"
            )

        keys = [p.key for p, _ in found]
        dupes = sorted({k for k in keys if keys.count(k) > 1})
        if dupes:
            raise LayoutError(
                f"{sheet_name}: duplicate period columns for {', '.join(dupes)}. "
                f"Refusing to choose between them."
            )

        # Verify the declared order rather than trusting it; a mis-declared order would
        # silently reverse every month-over-month comparison.
        ordinals = [p.ordinal() for p, _ in found]
        descending = all(a > b for a, b in zip(ordinals, ordinals[1:]))
        ascending = all(a < b for a, b in zip(ordinals, ordinals[1:]))
        if date_order == "newest_first" and not descending:
            raise LayoutError(
                f"{sheet_name}: config says date_order=newest_first but the headers "
                f"are not strictly newest-to-oldest: {keys}"
            )
        if date_order == "oldest_first" and not ascending:
            raise LayoutError(
                f"{sheet_name}: config says date_order=oldest_first but the headers "
                f"are not strictly oldest-to-newest: {keys}"
            )

        found.sort(key=lambda pair: pair[0].ordinal())
        return [p for p, _ in found], found

    def _row_label(self, sf, row: int, label_cols: list[int]) -> str | None:
        for col in label_cols:
            raw = sf.cell(row, col).value
            if isinstance(raw, str) and raw.strip():
                return " ".join(raw.replace("\xa0", " ").split())
            if raw is not None and not isinstance(raw, str):
                return str(raw).strip()
        return None

    # ---- structural tagging -------------------------------------------

    def _apply_income_statement_structure(self, sheet: Sheet, config: ClientConfig) -> None:
        """Tag rows with section / group / category using the config's anchors.

        Section membership is positional: the rows strictly between a section's header
        label and its total label. A line's name never decides its section.
        """
        rows_by_label: dict[str, list[int]] = {}
        for item in sheet.line_items:
            rows_by_label.setdefault(_norm(item.label), []).append(item.row)

        def anchor(label: str, what: str) -> int:
            matches = rows_by_label.get(_norm(label), [])
            if not matches:
                raise LayoutError(
                    f"{sheet.name}: {what} label {label!r} not found. "
                    f"Update config/clients/{config.client_id}.yaml to match the workbook."
                )
            if len(matches) > 1:
                raise LayoutError(
                    f"{sheet.name}: {what} label {label!r} appears on rows {matches}. "
                    f"Refusing to choose between them."
                )
            return matches[0]

        by_row = {item.row: item for item in sheet.line_items}

        for sec in config.sections:
            key = sec["key"]
            header_row = anchor(sec["header_label"], f"section {key!r} header")
            total_row = anchor(sec["total_label"], f"section {key!r} total")
            if total_row <= header_row:
                raise LayoutError(
                    f"{sheet.name}: section {key!r} total row {total_row} is not below "
                    f"its header row {header_row}"
                )

            if header_row in by_row:
                by_row[header_row].role = "header"
                by_row[header_row].section = key
            total_item = by_row[total_row]
            total_item.role = "section_total"
            total_item.section = key

            # Claim subtotal groups first so their members leave the direct pool.
            claimed: set[int] = set()
            for grp in sec.get("subtotal_groups") or []:
                g_header = anchor(grp["header_label"], f"group {grp['key']!r} header")
                g_total = anchor(grp["total_label"], f"group {grp['key']!r} total")
                if not (header_row < g_header < g_total < total_row):
                    raise LayoutError(
                        f"{sheet.name}: subtotal group {grp['key']!r} (rows "
                        f"{g_header}-{g_total}) does not sit inside section {key!r} "
                        f"(rows {header_row}-{total_row})"
                    )
                if g_header in by_row:
                    by_row[g_header].role = "header"
                    by_row[g_header].section = key
                    by_row[g_header].group = grp["name"]
                    by_row[g_header].group_key = grp["key"]
                for r in range(g_header + 1, g_total):
                    item = by_row.get(r)
                    if item is None:
                        continue
                    item.section = key
                    item.group = grp["name"]
                    item.group_key = grp["key"]
                    if item.role == "item":
                        item.category = config.categorise_opex(item.label)
                g_total_item = by_row[g_total]
                g_total_item.role = "group_total"
                g_total_item.section = key
                g_total_item.group = grp["name"]
                g_total_item.group_key = grp["key"]
                claimed.update(range(g_header, g_total + 1))

            for r in range(header_row + 1, total_row):
                if r in claimed:
                    continue
                item = by_row.get(r)
                if item is None:
                    continue
                item.section = key
                if item.role == "item" and key == "opex":
                    item.category = config.categorise_opex(item.label)

        for name, label in config.reported_lines.items():
            row = anchor(label, f"reported line {name!r}")
            by_row[row].role = "reported"
            by_row[row].section = by_row[row].section or None


def _coerce(raw) -> tuple[float | None, bool]:
    """Convert a cached cell value to float. Returns (value, was_unreadable).

    Blank stays `None`. A non-numeric, non-blank value is `None` and flagged — it is a
    data-quality problem for validation to report, not something to silently zero.
    """
    if raw is None:
        return None, False
    if isinstance(raw, bool):
        return None, True
    if isinstance(raw, (int, float)):
        return float(raw), False
    if isinstance(raw, str):
        text = raw.strip()
        if not text:
            return None, False
        cleaned = text.replace(",", "").replace("$", "").replace("\xa0", "")
        negative = cleaned.startswith("(") and cleaned.endswith(")")
        if negative:
            cleaned = cleaned[1:-1]
        try:
            value = float(cleaned)
        except ValueError:
            return None, True
        return (-value if negative else value), False
    return None, True


def _parse_period(raw, date_formats: list[str]) -> Period | None:
    if isinstance(raw, dt.datetime):
        return Period(raw.year, raw.month, raw.strftime("%b %Y"))
    if isinstance(raw, dt.date):
        return Period(raw.year, raw.month, raw.strftime("%b %Y"))
    text = " ".join(str(raw).replace("\xa0", " ").split())
    for fmt in date_formats:
        try:
            parsed = dt.datetime.strptime(text, fmt)
        except ValueError:
            continue
        return Period(parsed.year, parsed.month, text)
    return None
